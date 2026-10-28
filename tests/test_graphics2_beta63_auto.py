from __future__ import annotations

import pytest

from srstudio.graphics2.model import GraphicsDocument
from srstudio.graphics2.operations import GraphicsSession
from srstudio.graphics2.command_router import GraphicsCommandRouter
from srstudio.graphics2.promotion_layouts import plan_promotion
from srstudio.graphics2.item_slots import create_item_slot, bind_product_to_item_slot
from srstudio.graphics2.promotion_composition import apply_slot_composition
from srstudio.graphics2.auto_promotion import HighlightStore, propose_promotion


def products(count):
    return [{"id": f"SKU-{i}", "name": f"PRODUTO {i}", "price": f"{i+2},99", "unit": "UN"}
            for i in range(count)]


def test_highlights_follow_code_not_import_order(tmp_path):
    path = tmp_path / "highlights.json"
    store = HighlightStore(path)
    store.set("SKU-7", True)
    assert HighlightStore(path).contains("SKU-7")
    assert not HighlightStore(path).contains("SKU-8")
    store.set("SKU-7", False)
    assert not HighlightStore(path).contains("SKU-7")


@pytest.mark.parametrize("count", [24, 52])
def test_preview_is_read_only_and_confirmation_assigns_every_product(count):
    session = GraphicsSession(GraphicsDocument())
    session.document.metadata["products"] = products(count)
    router = GraphicsCommandRouter(session)
    before = session.document.to_dict()
    proposal = router.dispatch({"name": "preview_auto_promotion"})
    assert proposal.ok, proposal.message
    assert session.document.to_dict() == before
    assert sum(p["count"] for p in proposal.payload["pages"]) == count
    applied = router.dispatch({"name": "apply_auto_promotion"})
    assert applied.ok, applied.message
    assert len(session.document.pages) == 1 + len(proposal.payload["pages"])
    assigned = [s.product_id for p in session.document.pages[1:] for s in p.slots.values() if s.product_id]
    assert len(assigned) == count
    assert set(assigned) == {p["id"] for p in products(count)}
    assert router.dispatch({"name": "undo"}).ok
    assert session.document.to_dict() == before
    assert router.dispatch({"name": "redo"}).ok
    assert len(session.document.pages) == 1 + len(proposal.payload["pages"])


def test_starred_products_get_one_exclusive_page_each():
    proposal = propose_promotion(products(11), {"SKU-8", "SKU-9"})
    assert [(page.mode, page.product_ids, page.featured) for page in proposal.pages[:2]] == [
        ("individual", ("SKU-8",), True), ("individual", ("SKU-9",), True)]
    assert proposal.pages[2].count == 9
    assert not proposal.pages[2].featured


def test_24_products_with_three_stars_produce_three_exclusive_and_collective_pages():
    featured = {"SKU-1", "SKU-8", "SKU-19"}
    proposal = propose_promotion(products(24), featured)
    assert len(proposal.pages) == 5
    assert [page.product_ids for page in proposal.pages[:3]] == [
        ("SKU-1",), ("SKU-8",), ("SKU-19",)]
    assert all(page.featured for page in proposal.pages[:3])
    assert sum(page.count for page in proposal.pages[3:]) == 21
    assert not any(code in featured for page in proposal.pages[3:] for code in page.product_ids)
    assert len({code for page in proposal.pages for code in page.product_ids}) == 24


def test_confirmed_highlights_have_own_editable_pages_and_no_duplicates(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    for code in ("SKU-1", "SKU-8", "SKU-19"):
        HighlightStore().set(code, True)
    session = GraphicsSession(GraphicsDocument())
    session.document.metadata["products"] = products(24)
    router = GraphicsCommandRouter(session)
    before = session.document.to_dict()
    preview = router.dispatch({"name": "preview_auto_promotion"})
    assert preview.ok
    assert session.document.to_dict() == before
    assert [p["featured"] for p in preview.payload["pages"]] == [True, True, True, False, False]
    applied = router.dispatch({"name": "apply_auto_promotion"})
    assert applied.ok, applied.message
    pages = session.document.pages[1:]
    assert len(pages) == 5
    for page, code in zip(pages[:3], ("SKU-1", "SKU-8", "SKU-19")):
        filled = [slot for slot in page.slots.values() if slot.product_id]
        assert len(filled) == 1
        assert filled[0].product_id == code
        assert filled[0].metadata["preset_id"] == "individual-beta44"
        assert page.metadata["auto_promotion"]["featured"] is True
        assert page.height > page.width
        assert page.metadata.get("campaign_background_node")
        assert page.metadata.get("campaign_seal_node")
    collective = [slot.product_id for page in pages[3:] for slot in page.slots.values() if slot.product_id]
    assert len(collective) == 21
    assert not {"SKU-1", "SKU-8", "SKU-19"}.intersection(collective)
    assert router.dispatch({"name": "undo"}).ok
    assert session.document.to_dict() == before


def test_six_and_nine_match_expected_slot_families():
    assert [p.kind for p in plan_promotion("six").placements] == ["single"] * 6
    assert [p.kind for p in plan_promotion("nine").placements] == ["single"] * 4 + ["multi"] * 5


def test_six_migrates_filled_individual_into_real_single_and_hides_old_art():
    session = GraphicsSession(GraphicsDocument())
    old = create_item_slot(session, "individual-beta44")
    bind_product_to_item_slot(session, old.id, products(1)[0])
    result = apply_slot_composition(session, plan_promotion("six"))
    assert old.id not in result.slot_ids
    assert session.page.node(old.metadata["root_node_id"]).visible is False
    assert [session.page.slots[sid].metadata.get("limpeza_blue_profile") for sid in result.slot_ids] == ["limpeza-blue-v1"] * 6
    assert session.page.slots[result.slot_ids[0]].product_id == "SKU-0"


def test_custom_layout_is_editable_real_slot_topology():
    session = GraphicsSession(GraphicsDocument())
    plan = plan_promotion("custom")
    assert [p.kind for p in plan.placements] == ["individual"] + ["single"] * 3 + ["multi"] * 6
    result = apply_slot_composition(session, plan, reuse_existing=False)
    assert len(result.slot_ids) == 10
    assert all(session.page.slots[sid].metadata.get("root_node_id") for sid in result.slot_ids)
    assert max(p.y + p.height for p in plan.placements) <= .93


def test_nine_layout_moves_legacy_individuals_into_four_singles_and_real_multi():
    session = GraphicsSession(GraphicsDocument())
    old_roots = []
    for product in products(9):
        slot = create_item_slot(session, "individual-beta44")
        bind_product_to_item_slot(session, slot.id, product)
        old_roots.append(slot.metadata["root_node_id"])
    result = apply_slot_composition(session, plan_promotion("nine"))
    assert len(result.slot_ids) == 9
    assert [session.page.slots[sid].metadata.get("limpeza_blue_profile") for sid in result.slot_ids[:4]] == ["limpeza-blue-v1"] * 4
    assert [session.page.slots[sid].metadata.get("preset_id") for sid in result.slot_ids[4:]] == ["multi-cell-beta44"] * 5
    assert {session.page.slots[sid].product_id for sid in result.slot_ids} == {p["id"] for p in products(9)}
    assert all(not session.page.node(root).visible for root in old_roots)


def test_custom_layout_reapplication_keeps_real_mixed_families_with_filled_products():
    session = GraphicsSession(GraphicsDocument())
    old_roots = []
    for product in products(10):
        slot = create_item_slot(session, "individual-beta44")
        bind_product_to_item_slot(session, slot.id, product)
        old_roots.append(slot.metadata["root_node_id"])
    result = apply_slot_composition(session, plan_promotion("custom"))
    families = [session.page.slots[sid].metadata.get("preset_id") for sid in result.slot_ids]
    assert families[0] == "individual-beta44"
    assert families[4:] == ["multi-cell-beta44"] * 6
    assert {session.page.slots[sid].product_id for sid in result.slot_ids} == {p["id"] for p in products(10)}
    assert sum(session.page.node(root).visible for root in old_roots) == 1


def test_excel_reimport_keeps_highlights_by_product_code(tmp_path, monkeypatch):
    from openpyxl import Workbook
    monkeypatch.setenv("APPDATA", str(tmp_path))
    store = HighlightStore()
    store.set("1002", True)
    source = tmp_path / "terca-ofertas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["CÓDIGO", "PRODUTO", "PREÇO", "UNIDADE"])
    sheet.append(["1001", "ARROZ 5KG", 19.9, "UN"])
    sheet.append(["1002", "CAFÉ 500G", 12.99, "UN"])
    workbook.save(source)
    session = GraphicsSession(GraphicsDocument())
    session.import_product_catalog(str(source))
    proposal = GraphicsCommandRouter(session).dispatch({"name": "preview_auto_promotion"})
    assert proposal.ok, proposal.message
    first_id = proposal.payload["pages"][0]["product_ids"][0]
    assert next(p for p in session.document.metadata["products"] if p["id"] == first_id)["code"] == "1002"
    assert proposal.payload["campaign_id"] == "terca-verde"
    session.import_product_catalog(str(source))
    assert GraphicsCommandRouter(session).payload()["editor"]["highlighted_product_ids"] == ["1002"]


def test_auto_pages_remain_portrait_and_old_project_roundtrips():
    document = GraphicsDocument()
    document.active_page.width = 1350
    document.active_page.height = 1080
    document.metadata["products"] = products(24)
    session = GraphicsSession(GraphicsDocument.from_dict(document.to_dict()))
    result = GraphicsCommandRouter(session).dispatch({"name": "apply_auto_promotion"})
    assert result.ok, result.message
    restored = GraphicsDocument.from_dict(session.document.to_dict())
    assert len(restored.pages) == 4
    assert all(page.height > page.width for page in restored.pages[1:])
    assert restored.pages[0].width == 1350
    assert restored.pages[0].height == 1080


def test_auto_promotion_renders_png_and_pdf(tmp_path):
    pytest.importorskip("PySide6")
    from srstudio.graphics2.qt_renderer import render_pdf, render_png
    from pypdf import PdfReader
    session = GraphicsSession(GraphicsDocument())
    session.document.metadata["products"] = products(7)
    result = GraphicsCommandRouter(session).dispatch({"name": "apply_auto_promotion"})
    assert result.ok, result.message
    png = render_png(session.document, tmp_path / "montagem.png", page_index=1, target_width=540)
    pdf = render_pdf(session.document, tmp_path / "montagem.pdf", dpi=96)
    assert png.ok and png.width == 540 and png.output.stat().st_size > 1000
    assert pdf.ok and len(PdfReader(str(pdf.output)).pages) == 2


def test_preview_catalog_covers_all_variable_page_sizes():
    from srstudio.graphics2.slot_preview import PREVIEW_TYPES
    assert {"layout-auto-2", "layout-auto-3", "layout-auto-4", "layout-auto-5",
            "layout-nine-8", "layout-eleven-10"}.issubset(PREVIEW_TYPES)


@pytest.mark.parametrize("name,products,expected", [
    ("ofertas.xlsx", [
        {"id": "1", "name": "COXA SOBRECOXA 1KG", "category": "CARNES"},
        {"id": "2", "name": "ACÉM BOVINO KG"},
        {"id": "3", "name": "COSTELINHA SUÍNA KG"}], "quinta-file"),
    ("ofertas.xlsx", [
        {"id": "1", "name": "MAÇÃ NACIONAL GALA"},
        {"id": "2", "name": "MANGA PALMER"},
        {"id": "3", "name": "CENOURA KG"}], "terca-verde"),
    ("catalogo.xlsx", [
        {"id": "1", "name": "CAFÉ TORRADO 500G"},
        {"id": "2", "name": "CAFÉ EXTRA FORTE 500G"}], "quarta-cafe"),
    ("ofertas.xlsx", [
        {"id": "1", "name": "DETERGENTE 500ML"},
        {"id": "2", "name": "ARROZ BRANCO 5KG"}], "terca-verde"),
])
def test_campaign_is_selected_from_sheet_products_not_a_generic_filename(name, products, expected):
    from srstudio.graphics2.auto_promotion import propose_promotion
    proposal = propose_promotion(products, set(), source=name)
    assert proposal.campaign_id == expected
    assert proposal.background_id
    assert proposal.logo_id


def test_campaign_respects_previous_choice_when_products_are_ambiguous():
    from srstudio.graphics2.auto_promotion import propose_promotion
    sheet = [{"id": "1", "name": "COXA SOBRECOXA"},
             {"id": "2", "name": "MANGA PALMER"}]
    proposal = propose_promotion(sheet, set(), source="ofertas.xlsx",
                                 existing_campaign="churrasco")
    assert proposal.campaign_id == "churrasco"


def test_explicit_campaign_filename_overrides_mix_and_path_names():
    from srstudio.graphics2.auto_promotion import propose_promotion
    sheet = [{"id": "1", "name": "COXA SOBRECOXA"}]
    named = propose_promotion(sheet, set(),
                              source=r"C:\\Arquivos\\Quarta Café.xlsx")
    assert named.campaign_id == "quarta-cafe"
    generic = propose_promotion(sheet, set(),
                                source=r"C:\\Users\\lucas.cafe\\ofertas.xlsx")
    assert generic.campaign_id == "quinta-file"


def test_highlight_button_identifies_full_page_exclusive_flyer():
    from pathlib import Path
    qml = (Path(__file__).resolve().parents[1] /
           "src/srstudio/graphics2/qml/GraphicsEditor.qml").read_text(encoding="utf-8")
    assert 'ToolTip.text: "Criar um encarte exclusivo com este produto"' in qml
    assert "Priorizar nos Smart Slots individuais" not in qml


def test_explicit_terca_campaign_filename_wins_over_coffee_in_spreadsheet():
    from srstudio.graphics2.auto_promotion import propose_promotion
    sheet = [{"id": "1", "name": "CAFÉ TORRADO 500G"},
             {"id": "2", "name": "CAFÉ MOÍDO 500G"}]
    proposal = propose_promotion(sheet, set(), source="terca-ofertas.xlsx")
    assert proposal.campaign_id == "terca-verde"
